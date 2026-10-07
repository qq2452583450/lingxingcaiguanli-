import pytest
from openpyxl import load_workbook

from blueprints.exam import _exam_results_workbook


@pytest.mark.parametrize('include_history', [False, True])
def test_results_export_omits_subjective_column_and_preserves_scores(include_history):
    output = _exam_results_workbook([{
        'user_name': '测试考生',
        'username': 'test_user',
        'role_name': '材料员',
        'paper_title': '测试试卷',
        'status': 'completed',
        'objective_score': 76,
        'final_subjective_score': 17,
        'final_score': 93,
        'started_at': '2026-10-07 10:00:00',
        'submitted_at': '2026-10-07 10:30:00',
    }], include_history=include_history)
    sheet = load_workbook(output).active
    headers = [cell.value for cell in sheet[3]]
    assert headers == ['姓名', '账号', '角色', '试卷', '状态', '客观题', '总分', '开始时间', '提交时间', '备注']
    assert sheet['F4'].value == 76
    assert sheet['G4'].value == 93
    assert sheet['H4'].value == '2026-10-07 10:00:00'
    assert sheet['I4'].value == '2026-10-07 10:30:00'
    assert sheet.max_column == 10
    assert str(sheet.merged_cells) == 'A1:J1'
    assert sheet.auto_filter.ref == 'A3:J4'
