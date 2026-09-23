from io import BytesIO

from openpyxl import load_workbook


def _login(client):
    with client.session_transaction() as sess:
        sess['user'] = {
            'id': 1,
            'username': 'materials_exporter',
            'real_name': '材料导出员',
            'role_name': '材料员',
        }


def test_material_export_returns_all_matching_filtered_materials(client, test_db):
    cursor = test_db.cursor()
    cursor.execute("INSERT INTO units (unit_name) VALUES ('根')")
    unit_id = cursor.lastrowid
    cursor.execute("INSERT INTO suppliers (supplier_name) VALUES ('测试供应商')")
    supplier_id = cursor.lastrowid
    cursor.execute("INSERT INTO projects (project_code, project_name) VALUES ('QJLX', '曲靖项目')")
    project_id = cursor.lastrowid
    cursor.execute(
        """
        INSERT INTO materials (
            material_code, material_name, specification, detail_spec, brand, unit_id,
            tax_rate, tax_price, tax_exempt_price, cash_price, cash_tax_price,
            default_supplier_id, freight, inventory_min, inventory_max, weight,
            is_national_standard, project_id, create_time, remark
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        ('QJLX00001', '止水螺杆', 'Φ14', '墙厚250mm', '测试品牌', unit_id,
         0.13, 10, 8.85, 9, 7.96, supplier_id, 1.2, 3, 20, 0.5,
         1, project_id, '2026-09-23 10:00:00', '测试备注'),
    )
    cursor.execute(
        "INSERT INTO materials (material_code, material_name, specification) VALUES (?, ?, ?)",
        ('CDLX00001', '其他材料', 'DN20'),
    )
    test_db.commit()
    _login(client)

    response = client.get('/api/materials/export?filter_name=止水&filter_region=QJ')

    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ['导出说明', '材料信息']
    sheet = workbook['材料信息']
    assert sheet.max_row == 2
    assert [sheet.cell(1, index).value for index in range(1, 6)] == [
        '材料编号', '地区', '项目', '材料名称', '规格'
    ]
    assert [sheet.cell(2, index).value for index in range(1, 16)] == [
        'QJLX00001', '曲靖', '曲靖项目', '止水螺杆', 'Φ14', '墙厚250mm',
        '测试品牌', '根', '是', 0.13, 10, 8.85, 9, 7.96, '测试供应商',
    ]
    assert sheet['P2'].value == 1.2
    assert sheet['W2'].value == '测试备注'


def test_material_export_requires_login(client):
    response = client.get('/api/materials/export')

    assert response.get_json() == {'success': False, 'message': '未登录'}
